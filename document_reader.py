from pathlib import Path
from docx import Document
import pymupdf
import openpyxl
import xlrd
import subprocess
import tempfile
import shutil


def convert_doc_to_docx(file_path):
    """使用 LibreOffice 将 .doc 转换为 .docx。"""
    path = Path(file_path).resolve()
    output_dir = Path(tempfile.mkdtemp())

    soffice = shutil.which("soffice") or shutil.which("libreoffice")

    if not soffice:
        candidates = [
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
            Path(r"D:\program\soffice.exe"),
        ]
        soffice = next(
            (str(p) for p in candidates if p.exists()),
            None,
        )

    if not soffice:
        raise FileNotFoundError(
            "未找到 LibreOffice。请安装 LibreOffice，"
            "或在系统 PATH 中配置 soffice。"
        )

    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to", "docx",
            "--outdir", str(output_dir),
            str(path),
        ],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=60,
    )

    converted_path = output_dir / f"{path.stem}.docx"

    if not converted_path.exists():
        raise RuntimeError("DOC 转换失败，请检查文件是否有效。")

    return converted_path

def read_document(file_path):
    """读取不同格式的文档，返回文本。"""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".doc":
       path = convert_doc_to_docx(path)
       suffix = ".docx"

    # Word：.docx
    if suffix == ".docx":
        doc = Document(path)
        return "\n".join(
            p.text.strip()
            for p in doc.paragraphs
            if p.text.strip()
        )

    # PDF
    elif suffix == ".pdf":
        with pymupdf.open(path) as pdf:
            return "\n".join(page.get_text() for page in pdf)

    # Excel：.xlsx
    elif suffix == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(v) for v in row if v is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    # Excel：.xls
    elif suffix == ".xls":
        wb = xlrd.open_workbook(path)
        lines = []
        for sheet in wb.sheets():
            for row in range(sheet.nrows):
                values = [
                    str(v) for v in sheet.row_values(row)
                    if v != ""
                ]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)

    else:
        raise ValueError(f"暂不支持的文件格式：{suffix}")